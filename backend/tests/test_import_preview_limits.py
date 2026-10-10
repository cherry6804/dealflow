"""Regression tests for import preview row limits."""

from __future__ import annotations

import csv

from app.imports.preview import MAX_PREVIEW_ROWS, _preview_csv, _preview_xlsx
from openpyxl import Workbook


def test_csv_exact_preview_limit_has_no_more_rows(tmp_path) -> None:
    """Exactly 20 data rows should not indicate additional rows."""
    path = tmp_path / "exact.csv"

    with path.open("w", encoding="utf-8", newline="") as source:
        writer = csv.writer(source)
        writer.writerow(["First Name", "Email"])
        for index in range(MAX_PREVIEW_ROWS):
            writer.writerow([f"Contact {index}", f"contact{index}@example.com"])

    headers, rows, has_more_rows = _preview_csv(path)

    assert headers == ["First Name", "Email"]
    assert len(rows) == MAX_PREVIEW_ROWS
    assert has_more_rows is False


def test_csv_extra_nonempty_row_sets_has_more_rows(tmp_path) -> None:
    """A 21st non-empty data row should set has_more_rows."""
    path = tmp_path / "extra.csv"

    with path.open("w", encoding="utf-8", newline="") as source:
        writer = csv.writer(source)
        writer.writerow(["First Name", "Email"])
        for index in range(MAX_PREVIEW_ROWS + 1):
            writer.writerow([f"Contact {index}", f"contact{index}@example.com"])

    _, rows, has_more_rows = _preview_csv(path)

    assert len(rows) == MAX_PREVIEW_ROWS
    assert has_more_rows is True


def test_csv_blank_rows_do_not_count_as_more_data(tmp_path) -> None:
    """Blank rows after the preview limit should not set has_more_rows."""
    path = tmp_path / "blank_rows.csv"

    with path.open("w", encoding="utf-8", newline="") as source:
        writer = csv.writer(source)
        writer.writerow(["First Name", "Email"])
        for index in range(MAX_PREVIEW_ROWS):
            writer.writerow([f"Contact {index}", f"contact{index}@example.com"])
        writer.writerow([])
        writer.writerow(["", "  "])

    _, rows, has_more_rows = _preview_csv(path)

    assert len(rows) == MAX_PREVIEW_ROWS
    assert has_more_rows is False


def test_xlsx_exact_preview_limit_has_no_more_rows(tmp_path) -> None:
    """Exactly 20 non-empty XLSX rows should not indicate additional rows."""
    path = tmp_path / "exact.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Contacts"
    worksheet.append(["First Name", "Email"])

    for index in range(MAX_PREVIEW_ROWS):
        worksheet.append([f"Contact {index}", f"contact{index}@example.com"])

    workbook.save(path)
    workbook.close()

    sheet_name, headers, rows, has_more_rows = _preview_xlsx(path)

    assert sheet_name == "Contacts"
    assert headers == ["First Name", "Email"]
    assert len(rows) == MAX_PREVIEW_ROWS
    assert has_more_rows is False


def test_xlsx_extra_nonempty_row_sets_has_more_rows(tmp_path) -> None:
    """A 21st non-empty XLSX row should set has_more_rows."""
    path = tmp_path / "extra.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["First Name", "Email"])

    for index in range(MAX_PREVIEW_ROWS + 1):
        worksheet.append([f"Contact {index}", f"contact{index}@example.com"])

    workbook.save(path)
    workbook.close()

    _, _, rows, has_more_rows = _preview_xlsx(path)

    assert len(rows) == MAX_PREVIEW_ROWS
    assert has_more_rows is True


def test_xlsx_blank_rows_do_not_count_as_more_data(tmp_path) -> None:
    """Blank XLSX rows after the preview limit should not set has_more_rows."""
    path = tmp_path / "blank_rows.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["First Name", "Email"])

    for index in range(MAX_PREVIEW_ROWS):
        worksheet.append([f"Contact {index}", f"contact{index}@example.com"])

    worksheet.append([None, None])
    worksheet.append(["", "  "])

    workbook.save(path)
    workbook.close()

    _, _, rows, has_more_rows = _preview_xlsx(path)

    assert len(rows) == MAX_PREVIEW_ROWS
    assert has_more_rows is False


def test_csv_rejects_data_row_with_more_than_100_columns(tmp_path) -> None:
    """A sampled CSV row cannot exceed the supported column width."""
    from app.imports.preview import ImportPreviewFileError

    path = tmp_path / "too_many_columns.csv"

    with path.open("w", encoding="utf-8", newline="") as source:
        writer = csv.writer(source)
        writer.writerow(["First Name", "Email"])
        writer.writerow(["Contact", "contact@example.com"] + ["extra"] * 99)

    try:
        _preview_csv(path)
    except ImportPreviewFileError as exc:
        assert "columns" in str(exc).lower()
    else:
        raise AssertionError("Expected a row exceeding the column limit to fail")


def test_xlsx_rejects_sampled_data_in_column_101(tmp_path) -> None:
    """A populated 101st column in the sampled XLSX rows must be rejected."""
    from app.imports.preview import ImportPreviewFileError

    path = tmp_path / "too_many_columns.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["First Name", "Email"])
    row = ["Contact", "contact@example.com"] + [None] * 98 + ["extra"]
    worksheet.append(row)
    workbook.save(path)
    workbook.close()

    try:
        _preview_xlsx(path)
    except ImportPreviewFileError as exc:
        assert "columns" in str(exc).lower()
    else:
        raise AssertionError("Expected a row exceeding the column limit to fail")
