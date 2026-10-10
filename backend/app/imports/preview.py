"""Safe, read-only column detection and preview for import batches."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.models.import_batch import ImportBatch
from app.imports.storage import UploadValidationError, validate_source_file

MAX_PREVIEW_ROWS = 20
MAX_PREVIEW_COLUMNS = 100
MAX_PREVIEW_CELL_LENGTH = 1_000

CONTACT_FIELDS = {
    "first_name": "First name",
    "last_name": "Last name",
    "email": "Email",
    "phone": "Phone",
}

FIELD_ALIASES = {
    "first_name": {
        "first name",
        "firstname",
        "given name",
        "givenname",
        "contact first name",
    },
    "last_name": {
        "last name",
        "lastname",
        "surname",
        "family name",
        "familyname",
        "contact last name",
    },
    "email": {
        "email",
        "email address",
        "email id",
        "e-mail",
        "e-mail address",
    },
    "phone": {
        "phone",
        "phone number",
        "mobile",
        "mobile number",
        "telephone",
        "contact number",
    },
}


class ImportPreviewError(Exception):
    """Base exception for safe import preview failures."""


class ImportPreviewNotFoundError(ImportPreviewError):
    """The requested batch is unavailable to the current organization."""


class ImportPreviewFileError(ImportPreviewError):
    """The source file is missing or cannot be previewed."""


def _normalize_header(value: object) -> str:
    """Normalize a header for matching without changing its displayed value."""
    if value is None:
        return ""
    text = str(value).strip().casefold()
    return re.sub(r"[\s\_-]+", " ", text)


def _build_columns(headers: list[object]) -> list[dict[str, object]]:
    """Build stable, unique column descriptors and conservative suggestions."""
    columns: list[dict[str, object]] = []
    seen: dict[str, int] = {}

    for index, raw_header in enumerate(headers):
        original = "" if raw_header is None else str(raw_header).strip()
        base_label = original or f"Column {index + 1}"
        seen[base_label.casefold()] = seen.get(base_label.casefold(), 0) + 1
        occurrence = seen[base_label.casefold()]
        display_name = (
            base_label if occurrence == 1 else f"{base_label} ({occurrence})"
        )

        normalized = _normalize_header(original)
        suggested_field = None
        for field, aliases in FIELD_ALIASES.items():
            if normalized in aliases:
                suggested_field = field
                break

        columns.append(
            {
                "index": index,
                "source_name": original,
                "display_name": display_name,
                "suggested_field": suggested_field,
            }
        )

    return columns


def _safe_storage_path(batch: ImportBatch, organization_id: UUID) -> Path:
    """Resolve a local storage key strictly inside this tenant's storage area."""
    if batch.organization_id != organization_id:
        raise ImportPreviewNotFoundError("Import batch not found.")

    if batch.storage_backend != "local" or not batch.storage_key:
        raise ImportPreviewFileError(
            "The source file is not available in the configured storage backend."
        )

    settings = get_settings()
    root = settings.upload_storage_dir.resolve()
    candidate = (root / batch.storage_key).resolve()

    if candidate == root or root not in candidate.parents:
        raise ImportPreviewFileError("The stored source file reference is invalid.")

    # Stored keys are generated as <organization UUID>/<random filename>.
    if candidate.parent != (root / str(organization_id)).resolve():
        raise ImportPreviewFileError("The stored source file reference is invalid.")

    if not candidate.is_file():
        raise ImportPreviewFileError("The uploaded source file is no longer available.")

    return candidate


def _preview_csv(
    path: Path,
) -> tuple[list[object], list[list[object]], bool]:
    """Read a bounded CSV preview and detect whether more data rows exist."""
    try:
        with path.open("rb") as source:
            text = io.TextIOWrapper(
                source, encoding="utf-8-sig", errors="strict", newline=""
            )
            reader = csv.reader(text, strict=True)
            headers = next(reader, None)

            if not headers or not any(str(value).strip() for value in headers):
                raise ImportPreviewFileError(
                    "The CSV file does not contain a valid header row."
                )

            if len(headers) > MAX_PREVIEW_COLUMNS:
                raise ImportPreviewFileError(
                    f"The file has more than {MAX_PREVIEW_COLUMNS} columns."
                )

            rows: list[list[object]] = []
            has_more_rows = False


            for row in reader:
                if not any(str(value).strip() for value in row):
                    continue

                if len(row) > MAX_PREVIEW_COLUMNS:
                    raise ImportPreviewFileError(
                        f"The file has more than {MAX_PREVIEW_COLUMNS} columns."
                    )

                if len(rows) >= MAX_PREVIEW_ROWS:
                    has_more_rows = True
                    break

                rows.append(row)

            return headers, rows, has_more_rows
    except UnicodeDecodeError as exc:
        raise ImportPreviewFileError(
            "The CSV file must use UTF-8 encoding."
        ) from exc
    except csv.Error as exc:
        raise ImportPreviewFileError(
            "The CSV file contains malformed rows."
        ) from exc
    except OSError as exc:
        raise ImportPreviewFileError("The CSV file could not be read.") from exc


def _is_nonempty_cell(value: object) -> bool:
    """Return whether a cell has meaningful content."""
    return value is not None and (
        not isinstance(value, str) or bool(value.strip())
    )


def _preview_xlsx(
    path: Path,
) -> tuple[str, list[object], list[list[object]], bool]:
    """Read a bounded XLSX preview and detect additional non-empty rows."""
    workbook = None
    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        if not workbook.worksheets:
            raise ImportPreviewFileError("The workbook has no worksheets.")

        worksheet = workbook.worksheets[0]
        iterator = worksheet.iter_rows(
            min_row=1,
            max_col=MAX_PREVIEW_COLUMNS + 1,
            values_only=True,
        )

        first_row = next(iterator, None)
        if not first_row or not any(
            _is_nonempty_cell(value) for value in first_row
        ):
            raise ImportPreviewFileError(
                "The first worksheet does not contain a valid header row."
            )

        # Keep a bounded sample and skip blank rows.
        rows: list[list[object]] = []
        has_more_rows = False

        for row in iterator:
            if not any(_is_nonempty_cell(value) for value in row):
                continue

            if len(rows) >= MAX_PREVIEW_ROWS:
                has_more_rows = True
                break

            rows.append(list(row))

        # Determine the last meaningful column from headers and sampled data.
        # Blank cells in the middle remain intact; trailing empty columns go.
        last_nonempty_index = -1
        for index, value in enumerate(first_row):
            if _is_nonempty_cell(value):
                last_nonempty_index = index

        for row in rows:
            for index, value in enumerate(row):
                if _is_nonempty_cell(value):
                    last_nonempty_index = max(last_nonempty_index, index)

        if last_nonempty_index < 0:
            raise ImportPreviewFileError(
                "The first worksheet does not contain a valid header row."
            )

        if last_nonempty_index >= MAX_PREVIEW_COLUMNS:
            raise ImportPreviewFileError(
                f"The file has more than {MAX_PREVIEW_COLUMNS} columns."
            )

        column_count = last_nonempty_index + 1
        headers = list(first_row[:column_count])
        trimmed_rows = [row[:column_count] for row in rows]

        return worksheet.title, headers, trimmed_rows, has_more_rows
    except ImportPreviewFileError:
        raise
    except (OSError, ValueError, KeyError, EOFError) as exc:
        raise ImportPreviewFileError(
            "The Excel workbook could not be read."
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def _display_cell(value: object) -> str | None:
    """Return a bounded, JSON-safe display value."""
    if value is None:
        return None
    result = str(value)
    return result[:MAX_PREVIEW_CELL_LENGTH]


def get_import_preview(
    db: Session,
    *,
    organization_id: UUID,
    import_batch_id: UUID,
) -> dict[str, object]:
    """Return detected columns, mapping suggestions, and sample rows."""
    batch = db.scalar(
        select(ImportBatch).where(
            ImportBatch.id == import_batch_id,
            ImportBatch.organization_id == organization_id,
        )
    )
    if batch is None:
        raise ImportPreviewNotFoundError("Import batch not found.")

    path = _safe_storage_path(batch, organization_id)

    try:
        canonical_format = validate_source_file(path, batch.source_filename)
    except UploadValidationError as exc:
        raise ImportPreviewFileError(
            "The stored source file failed its format checks."
        ) from exc

    if canonical_format != batch.source_format:
        raise ImportPreviewFileError(
            "The stored file format does not match its import metadata."
        )

    worksheet_name = None
    if canonical_format == "csv":
        headers, raw_rows, has_more_rows = _preview_csv(path)
    elif canonical_format == "xlsx":
        (
            worksheet_name,
            headers,
            raw_rows,
            has_more_rows,
        ) = _preview_xlsx(path)
    else:
        raise ImportPreviewFileError("This source file format is not supported.")

    columns = _build_columns(headers)

    rows = [
        {
            str(column["display_name"]): _display_cell(
                raw_row[int(column["index"])]
                if int(column["index"]) < len(raw_row)
                else None
            )
            for column in columns
        }
        for raw_row in raw_rows
    ]

    return {
        "import_batch_id": batch.id,
        "source_filename": batch.source_filename,
        "source_format": canonical_format,
        "worksheet_name": worksheet_name,
        "columns": columns,
        "suggested_target_fields": [
            {"value": key, "label": label}
            for key, label in CONTACT_FIELDS.items()
        ],
        "preview_rows": rows,
        "preview_row_limit": MAX_PREVIEW_ROWS,
        "has_more_rows": has_more_rows,
        "contacts_created": 0,
    }
