"""Read-only full-file validation for contact imports."""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.db.models.import_batch import ImportBatch
from app.imports.preview import (
    CONTACT_FIELDS,
    FIELD_ALIASES,
    ImportPreviewFileError,
    ImportPreviewNotFoundError,
    _normalize_header,
    _safe_storage_path,
)
from app.imports.storage import UploadValidationError, validate_source_file
from app.imports.validation import validate_contact_rows

MAX_VALIDATION_ERRORS = 100
MAX_SOURCE_COLUMNS = 100
MAX_SOURCE_CELL_LENGTH = 10_000


class ImportValidationError(Exception):
    """Base exception for import validation failures."""


class ImportValidationNotFoundError(ImportValidationError):
    """The requested import batch is unavailable to this tenant."""


class ImportValidationFileError(ImportValidationError):
    """The stored import source cannot be safely validated."""


class ImportValidationMappingError(ImportValidationError):
    """The source columns cannot be mapped safely to contact fields."""


def _resolve_column_mapping(headers: list[object]) -> dict[str, int]:
    """Resolve only unambiguous header aliases to contact fields."""
    candidates: dict[str, list[int]] = {
        field: [] for field in CONTACT_FIELDS
    }

    for index, header in enumerate(headers):
        normalized = _normalize_header(header)
        for field, aliases in FIELD_ALIASES.items():
            if normalized in aliases:
                candidates[field].append(index)
                break

    ambiguous = [
        field for field, indices in candidates.items() if len(indices) > 1
    ]
    if ambiguous:
        fields = ", ".join(sorted(ambiguous))
        raise ImportValidationMappingError(
            f"Multiple source columns match these contact fields: {fields}. "
            "Resolve the column mapping before validating."
        )

    mapping = {
        field: indices[0]
        for field, indices in candidates.items()
        if indices
    }
    if "first_name" not in mapping:
        raise ImportValidationMappingError(
            "A source column matching First name is required. "
            "Use a supported first-name header before validating."
        )

    return mapping


def _validate_headers(headers: list[object]) -> None:
    """Reject missing, empty, or excessively wide source headers."""
    if not headers or not any(
        value is not None and str(value).strip() for value in headers
    ):
        raise ImportValidationFileError(
            "The source file does not contain a valid header row."
        )

    if len(headers) > MAX_SOURCE_COLUMNS:
        raise ImportValidationFileError(
            f"The source file has more than {MAX_SOURCE_COLUMNS} columns."
        )


def _normalize_source_row(
    raw_row: list[object],
    mapping: dict[str, int],
) -> dict[str, Any]:
    """Convert a source row into the supported contact field shape."""
    normalized: dict[str, Any] = {}
    for field, index in mapping.items():
        value = raw_row[index] if index < len(raw_row) else None
        if isinstance(value, str):
            value = value[:MAX_SOURCE_CELL_LENGTH]
        normalized[field] = value
    return normalized


def _iter_csv_rows(path: Path) -> Iterator[tuple[int, list[object]]]:
    """Yield non-empty CSV data rows with their physical source row numbers."""
    try:
        with path.open("rb") as source:
            text = io.TextIOWrapper(
                source,
                encoding="utf-8-sig",
                errors="strict",
                newline="",
            )
            reader = csv.reader(text, strict=True)
            headers = next(reader, None)
            _validate_headers(headers or [])

            for row_number, row in enumerate(reader, start=2):
                if not any(
                    value is not None and str(value).strip()
                    for value in row
                ):
                    continue
                if len(row) > MAX_SOURCE_COLUMNS:
                    raise ImportValidationFileError(
                        f"Row {row_number} exceeds the "
                        f"{MAX_SOURCE_COLUMNS}-column limit."
                    )
                yield row_number, row
    except ImportValidationFileError:
        raise
    except UnicodeDecodeError as exc:
        raise ImportValidationFileError(
            "The CSV file must use UTF-8 encoding."
        ) from exc
    except csv.Error as exc:
        raise ImportValidationFileError(
            "The CSV file contains malformed rows."
        ) from exc
    except OSError as exc:
        raise ImportValidationFileError(
            "The CSV file could not be read."
        ) from exc


def _read_csv_headers(path: Path) -> list[object]:
    """Read and validate the CSV header row."""
    try:
        with path.open("rb") as source:
            text = io.TextIOWrapper(
                source,
                encoding="utf-8-sig",
                errors="strict",
                newline="",
            )
            headers = next(csv.reader(text, strict=True), None)
            _validate_headers(headers or [])
            return headers or []
    except ImportValidationFileError:
        raise
    except UnicodeDecodeError as exc:
        raise ImportValidationFileError(
            "The CSV file must use UTF-8 encoding."
        ) from exc
    except (csv.Error, OSError) as exc:
        raise ImportValidationFileError(
            "The CSV header could not be read."
        ) from exc


def _iter_xlsx_rows(
    path: Path,
) -> Iterator[tuple[str, list[object], Iterator[tuple[int, list[object]]]]]:
    """Yield the first worksheet name, header row, and streaming data rows."""
    workbook = None
    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )
        if not workbook.worksheets:
            raise ImportValidationFileError(
                "The workbook has no worksheets."
            )

        worksheet = workbook.worksheets[0]
        rows = worksheet.iter_rows(
            min_row=1,
            max_col=MAX_SOURCE_COLUMNS + 1,
            values_only=True,
        )
        header_row = next(rows, None)
        headers = list(header_row or ())
        while headers and headers[-1] is None:
            headers.pop()
        _validate_headers(headers)
        if len(headers) > MAX_SOURCE_COLUMNS:
            raise ImportValidationFileError(
                f"The workbook has more than {MAX_SOURCE_COLUMNS} columns."
            )

        def data_rows() -> Iterator[tuple[int, list[object]]]:
            for row_number, values in enumerate(rows, start=2):
                row = list(values[: len(headers)])
                if not any(
                    value is not None
                    and (not isinstance(value, str) or value.strip())
                    for value in row
                ):
                    continue
                yield row_number, row

        yield worksheet.title, headers, data_rows()
    except ImportValidationFileError:
        raise
    except (OSError, ValueError, KeyError, EOFError) as exc:
        raise ImportValidationFileError(
            "The Excel workbook could not be read."
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def _collect_validation_result(
    rows: Iterator[tuple[int, list[object]]],
    mapping: dict[str, int],
) -> dict[str, Any]:
    """Validate rows incrementally while bounding retained error details."""
    total_rows = 0
    valid_rows = 0
    invalid_rows = 0
    errors: list[dict[str, Any]] = []
    errors_truncated = False

    for row_number, raw_row in rows:
        total_rows += 1
        normalized_row = _normalize_source_row(raw_row, mapping)
        result = validate_contact_rows([normalized_row])

        if result["invalid_rows"]:
            invalid_rows += 1
            for error in result["errors"]:
                error["row_number"] = row_number
                if len(errors) < MAX_VALIDATION_ERRORS:
                    errors.append(error)
                else:
                    errors_truncated = True
        else:
            valid_rows += 1

    return {
        "total_rows": total_rows,
        "valid_rows": valid_rows,
        "invalid_rows": invalid_rows,
        "errors": errors,
        "error_limit": MAX_VALIDATION_ERRORS,
        "errors_truncated": errors_truncated,
    }


def validate_import_source(
    db: Session,
    *,
    organization_id: UUID,
    import_batch_id: UUID,
) -> dict[str, Any]:
    """Validate an import source without creating contacts or changing state."""
    batch = db.scalar(
        select(ImportBatch).where(
            ImportBatch.id == import_batch_id,
            ImportBatch.organization_id == organization_id,
        )
    )
    if batch is None:
        raise ImportValidationNotFoundError("Import batch not found.")

    try:
        # Reuse the same tenant-safe local-storage boundary as preview.
        path = _safe_storage_path(batch, organization_id)
    except ImportPreviewNotFoundError as exc:
        raise ImportValidationNotFoundError(
            "Import batch not found."
        ) from exc
    except ImportPreviewFileError as exc:
        raise ImportValidationFileError(str(exc)) from exc

    try:
        canonical_format = validate_source_file(path, batch.source_filename)
    except UploadValidationError as exc:
        raise ImportValidationFileError(
            "The stored source file failed its format checks."
        ) from exc

    if canonical_format != batch.source_format:
        raise ImportValidationFileError(
            "The stored file format does not match its import metadata."
        )

    if canonical_format == "csv":
        headers = _read_csv_headers(path)
        mapping = _resolve_column_mapping(headers)
        rows = _iter_csv_rows(path)
        result = _collect_validation_result(rows, mapping)

    elif canonical_format == "xlsx":
        workbook_rows = _iter_xlsx_rows(path)
        try:
            worksheet_name, headers, rows = next(workbook_rows)
            mapping = _resolve_column_mapping(headers)
            result = _collect_validation_result(rows, mapping)
        finally:
            workbook_rows.close()
    else:
        raise ImportValidationFileError(
            "This source file format is not supported."
        )

    return {
        "import_batch_id": batch.id,
        "source_filename": batch.source_filename,
        "source_format": canonical_format,
        "worksheet_name": (
            worksheet_name if canonical_format == "xlsx" else None
        ),
        **result,
        "contacts_created": 0,
    }


__all__ = [
    "ImportValidationError",
    "ImportValidationFileError",
    "ImportValidationMappingError",
    "ImportValidationNotFoundError",
    "validate_import_source",
]
